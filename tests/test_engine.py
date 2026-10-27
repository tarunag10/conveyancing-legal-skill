"""Illustrative and adversarial legal scenario assertions, not professional sign-off."""
import json
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal

from conveyancing_skill.engine import (
    ROOT, LegalReviewRequired, completion, forms, money, load, review_matter,
    statutory_deadline, tax_estimate, workflow,
)
from conveyancing_skill.cli import validate, main
from conveyancing_skill.source_watch import is_official_url, watch
from conveyancing_skill.extract import extract_document

FX=ROOT/'tests'/'fixtures'

class RegistryTests(unittest.TestCase):
    def test_references_and_adapter_files_validate(self):
        res=validate(); self.assertEqual('PASS',res['status'],res['errors'])
        self.assertGreaterEqual(res['counts']['forms'],80)
        self.assertGreaterEqual(res['counts']['sources'],40)
        self.assertEqual(15,res['counts']['stages'])
        self.assertEqual(28,res['counts']['profiles'])

    def test_tr1_land_charges_and_licensed_forms_present(self):
        for name in ['TR1','TP1','FR1','DL','DS1','CH1','OC1','OC2','OS1','OS2','RX1','RX3','RX4','AN1','UN1','ID1','ID2','K15','TA6','TA7','TA10','TA13','LPE1','CPSE1']:
            with self.subTest(name=name):self.assertTrue(forms(name),f'Form missing: {name}')

    def test_no_form_claims_validated_binary_without_link(self):
        for form in load('knowledge/forms/catalogue.json'):
            with self.subTest(form=form['code']):
                self.assertFalse(form['file_revision_verified'])
                self.assertTrue(form['official_catalogue_url'].startswith('https://'))

    def test_proprietary_form_reproduction_not_claimed(self):
        for form in forms('TA6')+forms('LPE1')+forms('CPSE1'):
            self.assertIn('PROPRIETARY',form['copyright'])
            self.assertIsNone(form['official_download_link'])

    def test_all_sources_have_dates_not_approval(self):
        for source in load('knowledge/legislation/source-index.json'):
            self.assertTrue(source['last_verified_date'])
            self.assertTrue(source['scope_note'])
        for rule in load('knowledge/legislation/rules.json'):
            self.assertFalse(rule['legal_reviewer_approved'])

    def test_workflow_unproven_stages_and_profile(self):
        w=workflow('registered-freehold-purchase')
        self.assertEqual(15,len(w['stages']))
        self.assertTrue(all(s['default_status']=='outstanding' for s in w['stages']))
        self.assertTrue(w['profile']['human_approval_required'])
        with self.assertRaises(LegalReviewRequired):workflow('does-not-exist')

class FictionalCaseTests(unittest.TestCase):
    def check_case(self,filename,expected):
        j=json.loads((FX/(filename+'.json')).read_text())
        result=review_matter(j)
        self.assertEqual('NOT ASSESSED / NO AUTOMATED SIGN-OFF',result['completion_readiness'])
        codes=[issue['code'] for issue in result['issues']]
        for code in expected:self.assertIn(code,codes)
        return result

    def test_registered_freehold_no_fabricated_approval(self):
        result=self.check_case('registered-freehold',[])
        self.assertFalse(result['completion_blockers'])
        self.assertIn('NOT ASSESSED',result['completion_readiness'])

    def test_historic_charge_email_never_satisfies_discharge(self):
        result=self.check_case('historic-charge',['UNPROVEN_CHARGE_DISCHARGE'])
        self.assertIn('UNPROVEN_CHARGE_DISCHARGE',result['completion_blockers'])

    def test_auction_charge_and_missing_statement(self):
        self.check_case('auction-special',['OUTSTANDING_CHARGE','MISSING_DOCUMENT'])

    def test_leasehold_pack_missing(self):
        self.check_case('leasehold-management',['MISSING_DOCUMENT'])

    def test_ila_guarantee(self):
        self.check_case('mortgage-guarantee',['LENDER_CONDITION_OPEN'])

    def test_company_borrower(self):
        self.check_case('company-borrower',['LENDER_CONDITION_OPEN'])

    def test_adjoining_road_does_not_prove_adoption(self):
        result=self.check_case('public-highway',['HIGHWAY_STATUS_UNVERIFIED'])
        self.assertIn('HIGHWAY_STATUS_UNVERIFIED',result['completion_blockers'])

    def test_restrictive_covenants(self):
        self.check_case('restrictive-covenants',['UNSATISFIED_RESTRICTION'])

    def test_undisclosed_restriction(self):
        self.check_case('undisclosed-restriction',['UNSATISFIED_RESTRICTION'])

    def test_imminent_completion_lender_and_redemption(self):
        result=self.check_case('imminent-completion',['OUTSTANDING_CHARGE','LENDER_CONDITION_OPEN'])
        self.assertGreater(len(result['completion_blockers']),1)

    def test_commercial_togc(self):
        self.check_case('commercial-vat-togc',['MISSING_DOCUMENT'])

    def test_sdlt_and_ltt_require_effective_date_evidence(self):
        for x in ['sdlt','ltt']:
            with self.subTest(x=x):self.check_case(x,['MISSING_DOCUMENT'])

    def test_ap1_requisition_restriction(self):
        self.check_case('land-registry-requisition',['UNSATISFIED_RESTRICTION'])

    def test_conflicting_title_and_address(self):
        self.check_case('contradictory-pack',['TITLE_NUMBER_CONFLICT','ADDRESS_CONFLICT'])

    def test_ambiguous_stamp_is_not_discharge(self):
        self.check_case('ambiguous-discharge',['UNPROVEN_CHARGE_DISCHARGE'])

    def test_building_safety_issues(self):
        self.check_case('building-safety',['BUILDING_SAFETY_UNASSESSED','MISSING_DOCUMENT'])

    def test_bridging_special_conditions(self):
        self.check_case('bridging-finance',['LENDER_CONDITION_OPEN'])

    def test_seller_unsupported_claim_of_lender_waiver(self):
        m={'matter_id':'A','lender_requirements':[{'id':'a','status':'waived_by_lender','source_ref':'source','blocking':True}]}
        res=review_matter(m);self.assertIn('LENDER_CONDITION_OPEN',res['completion_blockers'])

    def test_seller_unsupported_verified_restriction(self):
        res=review_matter({'restrictions':[{'required':True,'status':'verified'}]})
        self.assertIn('UNSATISFIED_RESTRICTION',res['completion_blockers'])

    def test_positive_evidence_cannot_certify_completion(self):
        m={'charges':[{'id':'c','status':'discharged','evidence_type':'registered_discharge','evidence_ref':'OC1 page 3'}],
           'restrictions':[{'required':True,'status':'verified','evidence_ref':'certificate 25.10','checked_by_human':True}],
           'lender_requirements':[{'id':'L','status':'verified','source_ref':'actual lender ref','evidence_ref':'insurance ref','checked_by_human':True}]}
        res=review_matter(m)
        self.assertFalse(res['completion_blockers'])
        self.assertIn('NOT ASSESSED',res['completion_readiness'])

class MoneyTests(unittest.TestCase):
    def test_decimal_no_binary_float_arithmetic(self):
        self.assertEqual(Decimal('0.30'),money('0.30'))
        with self.assertRaises(LegalReviewRequired):money('NaN')
        with self.assertRaises(LegalReviewRequired):money(True)

    def test_completion_example_reconciliation_and_blockers(self):
        j=json.loads((ROOT/'examples/completion-input.json').read_text())
        r=completion(j)
        self.assertEqual('225100.75',r['balance_due_to_seller'])
        self.assertEqual('229300.75',r['total_required_at_completion'])
        self.assertEqual('-300.00',r['surplus_or_shortfall'])
        self.assertTrue(r['completion_blockers'])
        self.assertFalse(r['ready_to_send_funds'])

    def test_completion_not_autosend_even_if_inputs_confirm(self):
        r=completion({'side':'buyer','purchase_price':'100','deposit_paid':'10','cleared_client_funds':'90','cleared_lender_funds':'0','bank_evidence_confirmed_by_human':True,'lender_conditions_satisfied_by_human':True})
        self.assertEqual([],r['completion_blockers'])
        self.assertFalse(r['ready_to_send_funds'])

    def test_shortfall_and_malicious_negative_adjustments(self):
        with self.assertRaises(LegalReviewRequired):completion({'side':'buyer','purchase_price':100,'deposit_paid':20,'cleared_client_funds':50,'cleared_lender_funds':30,'seller_debits':[{'amount':-100}]})
        with self.assertRaises(LegalReviewRequired):completion({'side':'buyer','purchase_price':100,'deposit_paid':120,'cleared_client_funds':0,'cleared_lender_funds':0})

class TaxDeadlineTests(unittest.TestCase):
    def test_england_standard_as_of(self):
        r=tax_estimate('england','295000','2026-10-27',acknowledge_simple=True)
        self.assertEqual('4750.00',r['tax'])
        self.assertIn('ILLUSTRATIVE',r['legal_status'])

    def test_england_higher(self):
        r=tax_estimate('england',300000,'2026-10-27',True,acknowledge_simple=True)
        self.assertEqual('20000.00',r['tax'])

    def test_wales_standard(self):
        r=tax_estimate('wales','295000','2026-10-27',acknowledge_simple=True)
        self.assertEqual('4200.00',r['tax'])

    def test_wales_higher(self):
        r=tax_estimate('wales','250000','2026-10-27',True,acknowledge_simple=True)
        self.assertEqual('14950.00',r['tax'])

    def test_tax_refusals_no_reliefs_or_prospective_claims(self):
        with self.assertRaises(LegalReviewRequired):tax_estimate('england',300000,'2026-10-27')
        with self.assertRaises(LegalReviewRequired):tax_estimate('england',300000,'2026-10-28',acknowledge_simple=True)
        with self.assertRaises(LegalReviewRequired):tax_estimate('wales',300000,'2026-11-04',acknowledge_simple=True)
        with self.assertRaises(LegalReviewRequired):tax_estimate('wales',300000,'2026-10-27',acknowledge_simple=True,complexity_flags=['TOGC'])
        with self.assertRaises(LegalReviewRequired):tax_estimate('england',-12,'2026-10-27',acknowledge_simple=True)

    def test_sdlt_ltt_mr01_calendar_days_distinct(self):
        self.assertEqual('2026-11-10',statutory_deadline('sdlt','2026-10-27')['due_date'])
        self.assertEqual('2026-11-26',statutory_deadline('ltt','2026-10-27')['due_date'])
        self.assertEqual('2026-11-17',statutory_deadline('mr01','2026-10-27')['due_date'])

    def test_os_priority_requires_explicit_public_holidays(self):
        with self.assertRaises(LegalReviewRequired):statutory_deadline('os1','2026-10-27')
        r=statutory_deadline('os2','2026-10-27',holidays=['2026-12-25','2026-12-28'])
        self.assertEqual('working',r['day_type'])
        self.assertIn('Provisional',r['warning'])

    def test_no_generic_notice_deadline(self):
        with self.assertRaises(LegalReviewRequired):statutory_deadline('notice_to_complete','2026-10-27')

class ToolTests(unittest.TestCase):
    def test_cli_validate(self):self.assertEqual(0,main(['validate']))
    def test_extract_utf8_locators_and_warning(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'sample.txt';path.write_text('Land XX123456\n\nSigned by witness',encoding='utf8')
            d=extract_document(str(path))
            self.assertEqual('block 1',d['elements'][0]['location'])
            self.assertIn('XX123456',d['potential_title_numbers_NOT_VERIFIED'])
            self.assertTrue(any('authentic' in w for w in d['warnings']))

    def test_word_semantic_table_and_repeat_header(self):
        try:from docx import Document
        except ImportError:self.skipTest('python-docx optional dependency not installed')
        from conveyancing_skill.documents import to_docx
        with tempfile.TemporaryDirectory() as temp:
            s=Path(temp)/'a.md';o=Path(temp)/'a.docx'
            s.write_text('# Review\n\n| Title | Issue |\n|---|---|\n| AB12345 | **Restriction** |\n')
            to_docx(str(s),str(o))
            d=Document(str(o))
            self.assertEqual('Review',d.core_properties.title)
            self.assertEqual(1,len(d.tables))
            self.assertEqual('Title',d.tables[0].rows[0].cells[0].text)
            self.assertEqual('Restriction',d.tables[0].rows[1].cells[1].text)
            self.assertIn('tblHeader',d.tables[0].rows[0]._tr.xml)

    def test_empty_extraction_is_not_fabricated(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'blank.txt';path.write_text('',encoding='utf8')
            d=extract_document(str(path))
            self.assertEqual([],d['elements']);self.assertTrue(d['warnings'])

    def test_path_traversal_rejected(self):
        with self.assertRaises(LegalReviewRequired):load('../etc/passwd')

    def test_watch_default_no_network(self):
        r=watch(limit=2)
        self.assertFalse(r['network_enabled'])
        self.assertTrue(all(i['status']=='not_checked_network_disabled' for i in r['sources']))
        self.assertFalse(r['legal_review_queue'])

    def test_lender_candidate_extraction_is_never_cot_approval(self):
        from conveyancing_skill.intelligence import lender_candidates
        with tempfile.TemporaryDirectory() as tmp:
            f=Path(tmp)/'fictional-instructions.txt'
            f.write_text('1. Please provide building insurance policy before completion.\n\n2. Other information only.',encoding='utf8')
            r=lender_candidates(str(f))
            self.assertEqual(1,len(r['candidate_requirements']))
            self.assertEqual('EXTRACTED UNCHECKED',r['candidate_requirements'][0]['evidence_status'])
            self.assertEqual('UNDETERMINED',r['candidate_requirements'][0]['completion_blocker'])

    def test_cross_document_title_candidates_not_falsely_resolved(self):
        from conveyancing_skill.intelligence import compare_document_candidates
        with tempfile.TemporaryDirectory() as tmp:
            a=Path(tmp)/'a.txt';b=Path(tmp)/'b.txt'
            a.write_text('Title number AB123456, leasehold',encoding='utf8')
            b.write_text('Related freehold title XY654321',encoding='utf8')
            r=compare_document_candidates([str(a),str(b)])
            self.assertTrue(r['different_title_candidates_detected'])
            self.assertIn('legitimate',r['issues_to_check'][0])

    def test_official_domains_only(self):
        self.assertTrue(is_official_url('https://www.gov.uk/government/collections/hm-land-registry-forms'))
        for url in ('http://www.gov.uk/x','https://evil.gov.uk.attacker.test/x','https://127.0.0.1/secret','file:///etc/passwd','https://example.com'):
            with self.subTest(url=url):self.assertFalse(is_official_url(url))

    def test_docx_original_report_and_strict_replacement(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest('python-docx optional dependency not installed')
        from conveyancing_skill.documents import fill_placeholders,to_docx
        with tempfile.TemporaryDirectory() as temp:
            src=Path(temp)/'report.md';src.write_text('# Matter Report\n\nEvidence {{NUMBER}}',encoding='utf8')
            output=Path(temp)/'report.docx'
            to_docx(str(src),str(output))
            document=Document(str(output))
            self.assertIn('Matter Report',[p.text for p in document.paragraphs])
            original=Path(temp)/'original.docx';d=Document();d.add_paragraph('Title {{NUMBER}} unchanged paragraph.');d.add_paragraph('NEVER EDIT');d.save(str(original))
            modified=Path(temp)/'modified.docx'
            result=fill_placeholders(str(original),str(modified),{'NUMBER':'XX123456'},allowed=['NUMBER'],source_refs={'NUMBER':'verified official copy heading'})
            self.assertEqual(1,len(result['change_log']))
            new=Document(str(modified))
            self.assertEqual('Title XX123456 unchanged paragraph.',new.paragraphs[0].text)
            self.assertEqual('NEVER EDIT',new.paragraphs[1].text)
            with self.assertRaises(LegalReviewRequired):fill_placeholders(str(original),str(modified),{'NUMBER':'X'},allowed=['NUMBER'],source_refs={})
            with self.assertRaises(LegalReviewRequired):fill_placeholders(str(original),str(original),{'NUMBER':'X'},allowed=['NUMBER'],source_refs={'NUMBER':'a'})

if __name__=='__main__':unittest.main()
