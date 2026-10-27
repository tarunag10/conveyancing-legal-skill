"""Human-auditable extraction. OCR is never invoked silently.

Output is untrusted document evidence, not verified legal facts. Each paragraph is
located by page or paragraph index. Extracted signature text never verifies execution.
"""
from __future__ import annotations
import hashlib
import re
from pathlib import Path
from .engine import LegalReviewRequired


def extract_document(path: str, *, max_chars: int = 2_000_000) -> dict:
    p = Path(path)
    if not p.is_file():
        raise LegalReviewRequired(f'Document not found: {p}')
    suffix = p.suffix.lower()
    if suffix not in ('.txt', '.md', '.pdf', '.docx'):
        raise LegalReviewRequired('Unsupported type: use TXT, Markdown, PDF or DOCX')
    binary = p.read_bytes()
    if len(binary) > 40_000_000:
        raise LegalReviewRequired('Large files require controlled manual processing (40 MB limit)')
    doc_id = hashlib.sha256(binary).hexdigest()[:16]
    elements = []
    warnings = []
    if suffix == '.pdf':
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise LegalReviewRequired('PDF extraction requires pip install -e ".[documents]"') from exc
        try:
            reader=PdfReader(str(p))
            if reader.is_encrypted:
                raise LegalReviewRequired('Encrypted PDF: obtain authorised decrypted copy')
            for page_num, page in enumerate(reader.pages,1):
                page_text=page.extract_text(extraction_mode='plain') or ''
                if not page_text.strip():
                    warnings.append(f'PDF page {page_num} has no extractable text: manual inspection or authorised OCR required')
                for i,para in enumerate(re.split(r'\n\s*\n',page_text),1):
                    if para.strip(): elements.append({'location':f'page {page_num}, block {i}','text':para.strip()})
        except LegalReviewRequired:
            raise
        except Exception as exc:
            raise LegalReviewRequired(f'PDF extraction error: {exc}') from exc
    elif suffix == '.docx':
        try:
            from docx import Document
        except ImportError as exc:
            raise LegalReviewRequired('DOCX extraction requires pip install -e ".[documents]"') from exc
        try:
            document=Document(str(p))
            for i, paragraph in enumerate(document.paragraphs,1):
                if paragraph.text.strip():
                    elements.append({'location':f'paragraph {i}','text':paragraph.text.strip()})
            for ti, table in enumerate(document.tables,1):
                for ri,row in enumerate(table.rows,1):
                    for ci,cell in enumerate(row.cells,1):
                        if cell.text.strip():elements.append({'location':f'table {ti}, row {ri}, cell {ci}','text':cell.text.strip()})
            if document.sections:
                warnings.append('Headers, footers, text boxes, tracked changes and drawings must be checked in original Word file; extraction may omit them')
        except Exception as exc:
            raise LegalReviewRequired(f'DOCX extraction error: {exc}') from exc
    else:
        try:text=p.read_text(encoding='utf-8-sig')
        except UnicodeError as exc: raise LegalReviewRequired('Text document is not valid UTF-8') from exc
        for i,para in enumerate(re.split(r'\n\s*\n',text),1):
            if para.strip():elements.append({'location':f'block {i}','text':para.strip()})
    total=sum(len(x['text']) for x in elements)
    if total>max_chars:
        raise LegalReviewRequired('Extracted text exceeds safe size; split document for review')
    joined='\n'.join(x['text'] for x in elements)
    title_candidates=sorted(set(re.findall(r'\b[A-Z]{1,3}[0-9]{3,8}\b', joined)))
    if any(w in joined.lower() for w in ('signed by','in the presence of','witness')):
        warnings.append('Execution/signature wording detected; not proof of valid execution or authentic signatures')
    if not elements:warnings.append('No readable text extracted; do not infer document content')
    return {'file_name':p.name, 'document_id':doc_id,'format':suffix[1:], 'elements':elements,'potential_title_numbers_NOT_VERIFIED':title_candidates,'warnings':warnings,'legal_status':'EXTRACTED UNVERIFIED DATA; inspect original page/plan and confirm provenance'}
